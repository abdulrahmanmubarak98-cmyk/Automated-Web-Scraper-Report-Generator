from openpyxl import load_workbook
import pandas as pd
import logging
from config.settings import CLEANED_CSV, REPORT_XLSX
import os

logger = logging.getLogger("project")


def generate_excel_report(input_file=CLEANED_CSV):
    logger.info("Generating Excel report")

    df = pd.read_csv(input_file)
    df.to_excel(REPORT_XLSX, index=False)

    wb = load_workbook(REPORT_XLSX)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    price_col = headers.index("price") + 1

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=price_col).number_format = "£#,##0.00"

    wb.save(REPORT_XLSX)

    logger.info(f"Excel report generated: {REPORT_XLSX}")

    if os.path.exists(REPORT_XLSX):
        return REPORT_XLSX
    else:
        raise Exception("Excel report was not created successfully")
